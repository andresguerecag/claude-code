"""
Extrae la tabla de recetas (clave de platillo -> kg de cada insumo por unidad
vendida) directamente de las FORMULAS del Excel "Control diario" que ya usa
la taquería, en vez de hardcodear los numeros.

Por que asi: las formulas son la fuente de verdad del negocio (las escribio
la persona que sabe cuanta carne lleva cada platillo). Si el menu cambia,
basta con volver a correr este script sobre un archivo actualizado.

Uso:
    python3 extraer_recetas.py <ruta_control_diario.xlsx> [hoja] > recetas.json
"""
import json
import re
import sys

import openpyxl

INSUMOS_COLS = {
    "I": "carne",
    "J": "pastor",
    "K": "queso",
    "L": "tortillas",
    "M": "telera",
}

FORMULA_RE = re.compile(r"^=\+?H\d+\*([\d.]+)$")


def extraer_recetas(path: str, hoja: str | None = None) -> dict:
    wb = openpyxl.load_workbook(path, data_only=False)
    wb_val = openpyxl.load_workbook(path, data_only=True)

    if hoja is None:
        # usa la primera hoja que tenga encabezado "CLAVE" en A2 (formato conocido)
        for name in wb.sheetnames:
            if str(wb[name]["A2"].value or "").strip().upper() == "CLAVE":
                hoja = name
                break
    if hoja is None:
        raise ValueError("No encontre una hoja con encabezado 'CLAVE' en A2")

    ws = wb[hoja]
    ws_val = wb_val[hoja]

    recetas = {}
    fila = 3
    while True:
        clave_cell = ws_val.cell(row=fila, column=1).value
        nombre_cell = ws_val.cell(row=fila, column=2).value
        if clave_cell is None and nombre_cell is None:
            # dos filas vacias seguidas = fin de la tabla
            siguiente_clave = ws_val.cell(row=fila + 1, column=1).value
            siguiente_nombre = ws_val.cell(row=fila + 1, column=2).value
            if siguiente_clave is None and siguiente_nombre is None:
                break
            fila += 1
            continue
        if clave_cell:
            claves = [c.strip().upper() for c in str(clave_cell).replace(" ", "").split("/") if c.strip()]
            receta = {}
            for col_letter, insumo in INSUMOS_COLS.items():
                formula = ws[f"{col_letter}{fila}"].value
                if isinstance(formula, str):
                    m = FORMULA_RE.match(formula.replace(" ", ""))
                    if m:
                        receta[insumo] = float(m.group(1))
            if receta:
                for clave in claves:
                    recetas[clave] = {
                        "nombre": str(nombre_cell).strip() if nombre_cell else clave,
                        "receta_por_unidad": receta,
                    }
        fila += 1

    return recetas


if __name__ == "__main__":
    path = sys.argv[1]
    hoja = sys.argv[2] if len(sys.argv) > 2 else None
    recetas = extraer_recetas(path, hoja)
    print(json.dumps(recetas, indent=2, ensure_ascii=False))
