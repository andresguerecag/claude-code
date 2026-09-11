"""
Conciliacion de dinero (Fase 2): cuadre de caja diario por sucursal.

Todo se lee del mismo "Formato de corte" que ya suben para el inventario --
ese archivo ya trae, ademas del inventario, un cuadrito con ingresos
(efectivo, tarjeta, plataformas) y gastos del dia en efectivo.

Principio igual que en inventario: los conceptos de gasto vienen en texto
libre (ej. "SUPER", "EXTENCIONES") y hay que asignarlos a una categoria fija
(la misma que usa la familia en su Estado Financiero). Nunca se adivina esa
asignacion -- lo que no se ha categorizado a mano se muestra como pendiente,
igual que los platillos no identificados.
"""
from __future__ import annotations

import json
import re
from dataclasses import dataclass, field
from pathlib import Path

import openpyxl

from . import db

CATEGORIAS_PATH = Path(__file__).parent / "categorias_gasto.json"
MAPEO_GASTOS_PATH = Path(__file__).parent / "mapeo_gastos.json"
IGNORAR = "IGNORAR"

# Diferencia "normal" de caja, en pesos -- por debajo de esto no se alerta.
# Configurable: la familia mencionó que $1-2 pesos es demasiado estricto.
TOLERANCIA_SOBRA_FALTA = 20.0


def _normalizar(texto: str) -> str:
    return re.sub(r"\s+", " ", str(texto).strip().upper())


@dataclass
class GastoDia:
    concepto: str
    monto: float


@dataclass
class MatchGastos:
    categorizados: list = field(default_factory=list)   # (concepto_original, categoria, monto)
    sin_categorizar: list = field(default_factory=list)  # GastoDia
    ignorados: list = field(default_factory=list)        # GastoDia


# --- Categorias y mapeo (misma logica dual Postgres/local que recetas) ----

def _categorias_desde_json() -> list[dict]:
    with open(CATEGORIAS_PATH, encoding="utf-8") as f:
        return json.load(f)


def _sembrar_pg_si_vacio() -> None:
    con = db.conectar()
    with con, con.cursor() as cur:
        cur.execute("SELECT COUNT(*) FROM categorias_gasto")
        if cur.fetchone()[0] == 0:
            for cat in _categorias_desde_json():
                cur.execute(
                    "INSERT INTO categorias_gasto (nombre, grupo) VALUES (%s, %s) ON CONFLICT DO NOTHING",
                    (cat["nombre"], cat["grupo"]),
                )
        cur.execute("SELECT COUNT(*) FROM mapeo_gastos")
        if cur.fetchone()[0] == 0 and MAPEO_GASTOS_PATH.exists():
            with open(MAPEO_GASTOS_PATH, encoding="utf-8") as f:
                mapeo_inicial = json.load(f)
            for concepto, valor in mapeo_inicial.items():
                cur.execute(
                    "INSERT INTO mapeo_gastos (concepto, valor) VALUES (%s, %s) ON CONFLICT DO NOTHING",
                    (concepto, valor),
                )
    con.close()


def listar_categorias() -> list[dict]:
    if not db.usando_postgres():
        return _categorias_desde_json()
    db.inicializar_tablas()
    _sembrar_pg_si_vacio()
    con = db.conectar()
    with con.cursor() as cur:
        cur.execute("SELECT nombre, grupo FROM categorias_gasto ORDER BY grupo, nombre")
        filas = cur.fetchall()
    con.close()
    return [{"nombre": n, "grupo": g} for n, g in filas]


def guardar_categoria(nombre: str, grupo: str) -> None:
    nombre = nombre.strip().upper()
    grupo = grupo.strip() or "General"
    if not nombre:
        raise ValueError("El nombre de la categoria no puede estar vacio")

    if db.usando_postgres():
        db.inicializar_tablas()
        con = db.conectar()
        with con, con.cursor() as cur:
            cur.execute(
                """
                INSERT INTO categorias_gasto (nombre, grupo) VALUES (%s, %s)
                ON CONFLICT (nombre) DO UPDATE SET grupo = EXCLUDED.grupo
                """,
                (nombre, grupo),
            )
        con.close()
        return

    categorias = _categorias_desde_json()
    categorias = [c for c in categorias if c["nombre"] != nombre]
    categorias.append({"nombre": nombre, "grupo": grupo})
    with open(CATEGORIAS_PATH, "w", encoding="utf-8") as f:
        json.dump(categorias, f, indent=2, ensure_ascii=False)


def eliminar_categoria(nombre: str) -> None:
    nombre = nombre.strip().upper()
    if db.usando_postgres():
        db.inicializar_tablas()
        con = db.conectar()
        with con, con.cursor() as cur:
            cur.execute("DELETE FROM categorias_gasto WHERE nombre = %s", (nombre,))
        con.close()
        return

    categorias = [c for c in _categorias_desde_json() if c["nombre"] != nombre]
    with open(CATEGORIAS_PATH, "w", encoding="utf-8") as f:
        json.dump(categorias, f, indent=2, ensure_ascii=False)


def cargar_mapeo_gastos() -> dict:
    if not db.usando_postgres():
        if not MAPEO_GASTOS_PATH.exists():
            return {}
        with open(MAPEO_GASTOS_PATH, encoding="utf-8") as f:
            return json.load(f)

    db.inicializar_tablas()
    _sembrar_pg_si_vacio()
    con = db.conectar()
    with con.cursor() as cur:
        cur.execute("SELECT concepto, valor FROM mapeo_gastos")
        filas = cur.fetchall()
    con.close()
    return dict(filas)


def guardar_mapeo_gasto(concepto: str, categoria_o_ignorar: str) -> None:
    concepto_norm = _normalizar(concepto)

    if db.usando_postgres():
        db.inicializar_tablas()
        con = db.conectar()
        with con, con.cursor() as cur:
            cur.execute(
                """
                INSERT INTO mapeo_gastos (concepto, valor) VALUES (%s, %s)
                ON CONFLICT (concepto) DO UPDATE SET valor = EXCLUDED.valor
                """,
                (concepto_norm, categoria_o_ignorar),
            )
        con.close()
        return

    mapeo = cargar_mapeo_gastos()
    mapeo[concepto_norm] = categoria_o_ignorar
    with open(MAPEO_GASTOS_PATH, "w", encoding="utf-8") as f:
        json.dump(mapeo, f, indent=2, ensure_ascii=False)


def emparejar_gastos(gastos: list[GastoDia], mapeo: dict) -> MatchGastos:
    """A diferencia de las recetas, aqui NO se intenta adivinar por prefijos
    -- un concepto de gasto en texto libre solo se categoriza si alguien ya
    lo mapeo antes a mano. Todo lo demas se reporta como pendiente."""
    resultado = MatchGastos()
    for gasto in gastos:
        concepto_norm = _normalizar(gasto.concepto)
        decision = mapeo.get(concepto_norm)
        if decision == IGNORAR:
            resultado.ignorados.append(gasto)
        elif decision:
            resultado.categorizados.append((gasto.concepto, decision, gasto.monto))
        else:
            resultado.sin_categorizar.append(gasto)
    return resultado


# --- Lectura del cuadro de dinero dentro del Formato de corte -------------

def _buscar_valor_por_etiqueta(ws, etiquetas: list[str], col_min: int, col_max: int, fila_max: int) -> float | None:
    """Busca una celda de texto que coincida con alguna etiqueta (columnas
    col_min..col_max) y regresa el valor numerico de la celda inmediata a
    la derecha. Insensible a mayusculas/espacios."""
    objetivo = {_normalizar(e) for e in etiquetas}
    for row in ws.iter_rows(min_row=1, max_row=fila_max, min_col=col_min, max_col=col_max):
        for cell in row:
            if isinstance(cell.value, str) and _normalizar(cell.value) in objetivo:
                vecino = ws.cell(row=cell.row, column=cell.column + 1).value
                if isinstance(vecino, (int, float)):
                    return float(vecino)
    return None


def leer_dinero_formato_corte(path: str, hoja: str) -> dict:
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[hoja]

    # --- Gastos: fila con encabezados 'GASTOS' y 'MONTO' contiguos ---
    fila_gastos = col_concepto = col_monto = None
    for row in ws.iter_rows(min_row=1, max_row=20, min_col=8, max_col=13):
        for cell in row:
            if isinstance(cell.value, str) and _normalizar(cell.value) == "GASTOS":
                vecino = ws.cell(row=cell.row, column=cell.column + 1).value
                if isinstance(vecino, str) and _normalizar(vecino) == "MONTO":
                    fila_gastos, col_concepto, col_monto = cell.row, cell.column, cell.column + 1
        if fila_gastos:
            break

    gastos = []
    if fila_gastos:
        fila = fila_gastos + 1
        while fila < fila_gastos + 20:
            concepto = ws.cell(row=fila, column=col_concepto).value
            monto = ws.cell(row=fila, column=col_monto).value
            if isinstance(concepto, str) and _normalizar(concepto) == "TOTAL NOTAS":
                break
            if isinstance(concepto, str) and concepto.strip() and isinstance(monto, (int, float)):
                gastos.append(GastoDia(concepto=concepto.strip(), monto=float(monto)))
            fila += 1

    # --- Ingresos: etiquetas conocidas, columnas I a O (donde vive este cuadro) ---
    def buscar(*etiquetas):
        return _buscar_valor_por_etiqueta(ws, list(etiquetas), col_min=9, col_max=15, fila_max=35)

    efectivo = buscar("EFECTIVO")
    total_tarjeta = buscar("TOTAL TARJETA")
    total_plataforma = buscar("TOTAL PLATAFORMA")
    didi_tarjeta = buscar("DIDI TARJETA")
    rappi = buscar("rappi")
    uber = buscar("uber")
    debito = buscar("debito", "débito")
    credito = buscar("credito", "crédito")
    notas_tarj_efvo = buscar("notas+tarj+efvo")
    venta_tira = buscar("VENTA TIRA")
    sobra_falta = buscar("SOBRA+ o  FALTA-", "SOBRA+ o FALTA-", "SOBRA + O FALTA -")

    return {
        "efectivo": efectivo,
        "total_tarjeta": total_tarjeta,
        "total_plataforma": total_plataforma,
        "didi_tarjeta": didi_tarjeta,
        "rappi": rappi,
        "uber": uber,
        "debito": debito,
        "credito": credito,
        "notas_tarj_efvo": notas_tarj_efvo,
        "venta_tira": venta_tira,
        "sobra_falta": sobra_falta,
        "gastos": gastos,
    }


# --- Reporte consolidado ---------------------------------------------------

def leer_fecha_del_corte(path: str, hoja: str) -> str | None:
    """La fecha del corte vive en una celda de tipo fecha dentro del cuadro
    de dinero (sin etiqueta confiable al lado) -- se busca por tipo de dato,
    no por posicion fija."""
    import datetime as _dt

    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[hoja]
    for row in ws.iter_rows(min_row=1, max_row=15, max_col=20):
        for cell in row:
            if isinstance(cell.value, _dt.datetime):
                return cell.value.date().isoformat()
    return None


def generar_reporte_dinero(path_formato_corte: str, hoja_corte: str) -> dict:
    datos = leer_dinero_formato_corte(path_formato_corte, hoja_corte)
    if datos["efectivo"] is None and datos["venta_tira"] is None and not datos["gastos"]:
        raise ValueError(
            f"La hoja '{hoja_corte}' del formato de corte no tiene datos de dinero capturados todavia."
        )
    mapeo = cargar_mapeo_gastos()
    match = emparejar_gastos(datos["gastos"], mapeo)

    total_gastos_categorizados = sum(m for _, _, m in match.categorizados)
    total_sin_categorizar = sum(g.monto for g in match.sin_categorizar)
    total_ignorado = sum(g.monto for g in match.ignorados)
    total_gastos = total_gastos_categorizados + total_sin_categorizar + total_ignorado

    gastos_por_categoria: dict[str, float] = {}
    for _concepto, categoria, monto in match.categorizados:
        gastos_por_categoria[categoria] = gastos_por_categoria.get(categoria, 0.0) + monto

    sobra_falta = datos["sobra_falta"]
    alerta_caja = abs(sobra_falta) > TOLERANCIA_SOBRA_FALTA if sobra_falta is not None else None

    total_gastos_relevante = total_gastos_categorizados + total_sin_categorizar
    pct_categorizado = (
        round(100 * total_gastos_categorizados / total_gastos_relevante, 1)
        if total_gastos_relevante else 100.0
    )

    return {
        "ingresos": {
            "efectivo": datos["efectivo"],
            "total_tarjeta": datos["total_tarjeta"],
            "total_plataforma": datos["total_plataforma"],
            "desglose_plataforma": {
                "didi_tarjeta": datos["didi_tarjeta"],
                "rappi": datos["rappi"],
                "uber": datos["uber"],
            },
            "desglose_tarjeta": {
                "debito": datos["debito"],
                "credito": datos["credito"],
            },
            "venta_declarada": datos["notas_tarj_efvo"],
            "venta_sistema": datos["venta_tira"],
        },
        "sobra_falta": sobra_falta,
        "alerta_caja": alerta_caja,
        "gastos_por_categoria": [
            {"categoria": cat, "monto": round(monto, 2)} for cat, monto in sorted(gastos_por_categoria.items())
        ],
        "gastos_sin_categorizar": [
            {"concepto": g.concepto, "monto": g.monto} for g in match.sin_categorizar
        ],
        "gastos_ignorados": [
            {"concepto": g.concepto, "monto": g.monto} for g in match.ignorados
        ],
        "total_gastos": round(total_gastos, 2),
        "total_gastos_categorizados": round(total_gastos_categorizados, 2),
        "total_gastos_sin_categorizar": round(total_sin_categorizar, 2),
        "pct_categorizado": pct_categorizado,
        "categorias_disponibles": [c["nombre"] for c in listar_categorias()],
    }
