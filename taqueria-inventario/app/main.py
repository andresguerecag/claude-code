"""
App web sencilla: arrastras el "Formato de corte" (excel que manda la
sucursal cada noche) y el reporte de ventas de Wansoft, y te regresa el
reporte de conciliacion de inventario del dia.
"""
import logging
import os
import re
import secrets
import tempfile
from pathlib import Path

from fastapi import Body, Depends, FastAPI, File, Form, HTTPException, Request, UploadFile
from fastapi.responses import FileResponse, JSONResponse, StreamingResponse
from fastapi.security import HTTPBasic, HTTPBasicCredentials
from fastapi.staticfiles import StaticFiles

_logger = logging.getLogger("uvicorn.error")

from . import asesor_compras, colchon, compras, dinero, exportar, historial, historial_dinero, reconciliacion, resumen_ia

STATIC_DIR = Path(__file__).parent / "static"

# Contraseña compartida para proteger la app cuando esté en internet. Se lee
# de variables de entorno (nunca escrita en el código) -- si no se configura
# ninguna, la app queda SIN contraseña (útil para correrla local en tu compu
# sin que te pida usuario/clave cada vez).
_APP_USUARIO = os.environ.get("APP_USUARIO")
_APP_PASSWORD = os.environ.get("APP_PASSWORD")
_seguridad = HTTPBasic(auto_error=False)


def requiere_acceso(credenciales: HTTPBasicCredentials | None = Depends(_seguridad)):
    if not _APP_USUARIO or not _APP_PASSWORD:
        return  # sin variables de entorno configuradas -> app abierta (uso local)
    ok_usuario = credenciales is not None and secrets.compare_digest(credenciales.username, _APP_USUARIO)
    ok_clave = credenciales is not None and secrets.compare_digest(credenciales.password, _APP_PASSWORD)
    if not (ok_usuario and ok_clave):
        raise HTTPException(status_code=401, detail="Acceso no autorizado", headers={"WWW-Authenticate": "Basic"})


app = FastAPI(title="Conciliacion de inventario - Taqueria", dependencies=[Depends(requiere_acceso)])


@app.exception_handler(Exception)
async def manejar_error_inesperado(request: Request, exc: Exception):
    """Si algo truena sin que lo hayamos previsto, esto evita que el
    navegador reciba una pagina de error en HTML (que no se puede leer como
    JSON y confunde con un mensaje tipo "Unexpected token") -- en vez de
    eso regresa el mensaje real del error, para poder diagnosticarlo."""
    _logger.exception("Error inesperado en %s", request.url.path)
    return JSONResponse(
        status_code=500,
        content={"detail": f"Error inesperado del servidor ({type(exc).__name__}): {exc}"},
    )


def _detectar_metadatos_wansoft(path_wansoft: str) -> dict:
    """
    Lee del propio reporte de Wansoft (sin que la usuaria tenga que
    tecleerlo):
      - la fecha completa ('Reporte del: 2026-09-05 al 2026-09-05')
      - el nombre de la sucursal ('Sucursal: T-Grill Sucursal Santa Catarina')
    """
    import openpyxl

    wb = openpyxl.load_workbook(path_wansoft, data_only=True)
    ws = wb[wb.sheetnames[0]]
    fecha_completa = dia = sucursal = None
    for row in ws.iter_rows(min_row=1, max_row=10):
        for cell in row:
            if not isinstance(cell.value, str):
                continue
            texto = cell.value.strip()
            if fecha_completa is None and "reporte del" in texto.lower():
                m = re.search(r"(\d{4}-\d{2}-\d{2})", texto)
                if m:
                    fecha_completa = m.group(1)
                    dia = fecha_completa.split("-")[2]
            if sucursal is None and texto.lower().startswith("sucursal:"):
                sucursal = texto.split(":", 1)[1].strip()
    return {"fecha": fecha_completa, "dia": dia, "sucursal": sucursal}


def _resolver_fecha_dia_sucursal(path_wansoft: str, fecha_elegida: str | None) -> dict:
    """Decide que dia/hoja usar. Si la usuaria elige una fecha a mano (nuevo
    selector), esa manda -- pero si el propio Wansoft trae una fecha
    distinta, se avisa en vez de proceder callado (podria ser que subio el
    archivo de otro dia por error)."""
    meta = _detectar_metadatos_wansoft(path_wansoft)
    if meta["sucursal"] is None:
        raise HTTPException(
            status_code=400,
            detail="No pude encontrar la sucursal ('Sucursal: ...') en el archivo que subiste como Wansoft. "
            "Verifica que sea el reporte 'Ventas Por Platillo Por Grupo' que exporta Wansoft "
            "(no otro archivo, como el control diario que se lleva aparte).",
        )

    if fecha_elegida:
        dia = fecha_elegida.split("-")[-1]
        if meta["dia"] and meta["dia"] != dia:
            raise HTTPException(
                status_code=400,
                detail=f"Elegiste el dia {dia}, pero el reporte de Wansoft dice que es del dia {meta['dia']}. "
                "Verifica que subiste el Wansoft del mismo dia que elegiste.",
            )
        return {"dia": dia, "fecha": fecha_elegida, "sucursal": meta["sucursal"], "origen": "La fecha que elegiste"}

    if meta["dia"] is None:
        raise HTTPException(
            status_code=400,
            detail="No pude encontrar la fecha ('Reporte del: ...') en el archivo de Wansoft. "
            "Verifica que sea el reporte de Ventas Por Platillo Por Grupo sin editar, o elige la fecha a mano.",
        )
    return {"dia": meta["dia"], "fecha": meta["fecha"], "sucursal": meta["sucursal"], "origen": "El reporte de Wansoft"}


def _elegir_hoja(path_formato_corte: str, dia: str, origen_fecha: str = "la fecha indicada") -> str:
    import openpyxl

    wb = openpyxl.load_workbook(path_formato_corte, read_only=True)
    for candidato in (dia, dia.lstrip("0") or "0", dia.zfill(2)):
        if candidato in wb.sheetnames:
            return candidato
    raise HTTPException(
        status_code=400,
        detail=f"{origen_fecha} es del dia {dia}, pero el formato de corte no tiene una hoja para ese dia. "
        "Verifica que subiste el 'Formato de corte' correcto (no el 'Formato de efectivo').",
    )


@app.post("/api/reconciliar")
async def reconciliar(
    formato_corte: UploadFile = File(...),
    wansoft: UploadFile = File(...),
    fecha: str | None = Form(None),
):
    with tempfile.TemporaryDirectory() as tmp:
        path_corte = Path(tmp) / "formato_corte.xlsx"
        path_wansoft = Path(tmp) / "wansoft.xlsx"
        path_corte.write_bytes(await formato_corte.read())
        path_wansoft.write_bytes(await wansoft.read())

        info = _resolver_fecha_dia_sucursal(str(path_wansoft), fecha)
        hoja = _elegir_hoja(str(path_corte), info["dia"], origen_fecha=info["origen"])

        try:
            reporte = reconciliacion.generar_reporte(
                path_formato_corte=str(path_corte),
                hoja_corte=hoja,
                path_wansoft=str(path_wansoft),
            )
        except ValueError as e:
            raise HTTPException(status_code=400, detail=str(e))

        reporte["dia_detectado"] = info["dia"]
        reporte["fecha"] = info["fecha"]
        reporte["sucursal"] = info["sucursal"]
        historial.guardar_reporte(info["fecha"], info["sucursal"], reporte)
        return reporte


@app.get("/api/historial")
async def ver_historial(sucursal: str | None = None, desde: str | None = None, hasta: str | None = None):
    return historial.listar_historial(sucursal=sucursal, desde=desde, hasta=hasta)


@app.delete("/api/historial/{fecha}/{sucursal}")
async def borrar_reporte_historial(fecha: str, sucursal: str):
    historial.eliminar_reporte(fecha, sucursal)
    return {"ok": True}


@app.get("/api/dashboard")
async def ver_dashboard(sucursal: str | None = None, desde: str | None = None, hasta: str | None = None):
    return historial.dashboard(sucursal=sucursal, desde=desde, hasta=hasta)


@app.get("/api/mermas")
async def ver_mermas(sucursal: str | None = None, desde: str | None = None, hasta: str | None = None, insumo: str | None = None):
    return historial.mermas_detalle(sucursal=sucursal, desde=desde, hasta=hasta, insumo=insumo)


@app.get("/api/resumen-ia/disponible")
async def resumen_ia_disponible():
    return {"disponible": resumen_ia.disponible()}


@app.get("/api/resumen-ia")
async def ver_resumen_ia(sucursal: str | None = None, desde: str | None = None, hasta: str | None = None):
    datos = historial.dashboard(sucursal=sucursal, desde=desde, hasta=hasta)
    texto = resumen_ia.generar_resumen(datos)
    if texto is None:
        raise HTTPException(status_code=503, detail="El resumen con IA no está configurado (falta ANTHROPIC_API_KEY).")
    return {"resumen": texto}


@app.get("/api/pendientes-acumulados")
async def ver_pendientes_acumulados():
    """Platillos no identificados de todo el historial, sumados y ordenados
    por volumen -- para priorizar una sesión de afinar recetas."""
    return historial.pendientes_acumulados()


@app.get("/api/historial/{fecha}/{sucursal}")
async def ver_reporte_historial(fecha: str, sucursal: str):
    reporte = historial.obtener_reporte(fecha, sucursal)
    if reporte is None:
        raise HTTPException(status_code=404, detail="No hay reporte guardado para esa fecha y sucursal.")
    return reporte


@app.post("/api/exportar")
async def exportar_reporte(
    formato_corte: UploadFile = File(...),
    wansoft: UploadFile = File(...),
    fecha: str | None = Form(None),
):
    """Regenera el mismo reporte y lo entrega como Excel descargable."""
    with tempfile.TemporaryDirectory() as tmp:
        path_corte = Path(tmp) / "formato_corte.xlsx"
        path_wansoft = Path(tmp) / "wansoft.xlsx"
        path_corte.write_bytes(await formato_corte.read())
        path_wansoft.write_bytes(await wansoft.read())

        info = _resolver_fecha_dia_sucursal(str(path_wansoft), fecha)
        hoja = _elegir_hoja(str(path_corte), info["dia"], origen_fecha=info["origen"])

        try:
            reporte = reconciliacion.generar_reporte(
                path_formato_corte=str(path_corte), hoja_corte=hoja, path_wansoft=str(path_wansoft),
            )
        except ValueError as e:
            raise HTTPException(status_code=400, detail=str(e))
        reporte["dia_detectado"] = info["dia"]
        reporte["fecha"] = info["fecha"]
        reporte["sucursal"] = info["sucursal"]

    buffer = exportar.generar_excel_reporte(reporte)
    nombre_archivo = f"conciliacion_{info['fecha']}_{info['sucursal'].replace(' ', '-')}.xlsx"
    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{nombre_archivo}"'},
    )


@app.post("/api/mapeo")
async def guardar_mapeo(payload: dict = Body(...)):
    """Guarda la correccion manual de una clave no identificada: a que clave
    de receta corresponde, o IGNORAR si no debe entrar al calculo (ej.
    refrescos, que se comparan distinto)."""
    clave = payload.get("clave")
    receta_o_ignorar = payload.get("receta_o_ignorar")
    if not clave or not receta_o_ignorar:
        raise HTTPException(status_code=400, detail="Falta 'clave' o 'receta_o_ignorar'.")
    reconciliacion.guardar_mapeo_manual(clave, receta_o_ignorar)
    return {"ok": True}


@app.get("/api/mapeo")
async def listar_mapeo():
    """Todos los mapeos manuales guardados, para poder revisarlos y
    corregirlos (ej. algo marcado como 'ignorar' que ahora si se quiere
    contemplar)."""
    return reconciliacion.listar_mapeo_manual()


@app.delete("/api/mapeo/{clave}")
async def eliminar_mapeo(clave: str):
    reconciliacion.eliminar_mapeo_manual(clave)
    return {"ok": True}


@app.get("/api/recetas")
async def listar_recetas():
    return reconciliacion.listar_recetas()


@app.post("/api/recetas")
async def guardar_receta(payload: dict = Body(...)):
    clave = payload.get("clave")
    nombre = payload.get("nombre", "")
    receta_por_unidad = payload.get("receta_por_unidad", {})
    if not clave:
        raise HTTPException(status_code=400, detail="Falta 'clave'.")
    try:
        reconciliacion.guardar_receta(clave, nombre, receta_por_unidad)
    except ValueError as e:
        raise HTTPException(status_code=400, detail=str(e))
    return {"ok": True}


@app.delete("/api/recetas/{clave}")
async def eliminar_receta(clave: str):
    reconciliacion.eliminar_receta(clave)
    return {"ok": True}


def _leer_fecha_hoja_corte(path_corte: str, hoja: str) -> str | None:
    return dinero.leer_fecha_del_corte(path_corte, hoja)


@app.post("/api/dinero/reconciliar")
async def reconciliar_dinero(
    formato_corte: UploadFile = File(...),
    sucursal: str = Form(...),
    fecha: str = Form(...),
):
    """A diferencia del inventario, aqui solo se sube el Formato de corte
    (no hace falta Wansoft) -- pero ese archivo no trae el nombre de la
    sucursal, asi que se pide en el formulario. La fecha exacta se toma del
    propio archivo cuando esta disponible (mas confiable que lo que se
    escriba a mano)."""
    if not sucursal.strip():
        raise HTTPException(status_code=400, detail="Falta indicar la sucursal.")
    with tempfile.TemporaryDirectory() as tmp:
        path_corte = Path(tmp) / "formato_corte.xlsx"
        path_corte.write_bytes(await formato_corte.read())

        dia = fecha.split("-")[-1] if fecha else ""
        if not dia:
            raise HTTPException(status_code=400, detail="Falta indicar la fecha.")
        hoja = _elegir_hoja(str(path_corte), dia, origen_fecha="La fecha que indicaste")

        try:
            reporte = dinero.generar_reporte_dinero(str(path_corte), hoja)
        except ValueError as e:
            raise HTTPException(status_code=400, detail=str(e))

        fecha_detectada = _leer_fecha_hoja_corte(str(path_corte), hoja) or fecha
        reporte["fecha"] = fecha_detectada
        reporte["sucursal"] = sucursal.strip()
        historial_dinero.guardar_reporte(fecha_detectada, sucursal.strip(), reporte)
        return reporte


@app.get("/api/dinero/historial")
async def ver_historial_dinero(sucursal: str | None = None, desde: str | None = None, hasta: str | None = None):
    return historial_dinero.listar_historial(sucursal=sucursal, desde=desde, hasta=hasta)


@app.get("/api/dinero/historial/{fecha}/{sucursal}")
async def ver_reporte_historial_dinero(fecha: str, sucursal: str):
    reporte = historial_dinero.obtener_reporte(fecha, sucursal)
    if reporte is None:
        raise HTTPException(status_code=404, detail="No hay reporte de dinero guardado para esa fecha y sucursal.")
    return reporte


@app.delete("/api/dinero/historial/{fecha}/{sucursal}")
async def borrar_reporte_historial_dinero(fecha: str, sucursal: str):
    historial_dinero.eliminar_reporte(fecha, sucursal)
    return {"ok": True}


@app.get("/api/dinero/dashboard")
async def ver_dashboard_dinero(sucursal: str | None = None, desde: str | None = None, hasta: str | None = None):
    return historial_dinero.dashboard(sucursal=sucursal, desde=desde, hasta=hasta)


@app.get("/api/dinero/pendientes-acumulados")
async def ver_pendientes_acumulados_dinero():
    return historial_dinero.pendientes_acumulados()


@app.get("/api/dinero/categorias")
async def listar_categorias_gasto():
    return dinero.listar_categorias()


@app.post("/api/dinero/categorias")
async def guardar_categoria_gasto(payload: dict = Body(...)):
    nombre = payload.get("nombre")
    grupo = payload.get("grupo", "General")
    if not nombre:
        raise HTTPException(status_code=400, detail="Falta 'nombre'.")
    try:
        dinero.guardar_categoria(nombre, grupo)
    except ValueError as e:
        raise HTTPException(status_code=400, detail=str(e))
    return {"ok": True}


@app.delete("/api/dinero/categorias/{nombre}")
async def eliminar_categoria_gasto(nombre: str):
    dinero.eliminar_categoria(nombre)
    return {"ok": True}


@app.post("/api/dinero/mapeo")
async def guardar_mapeo_gasto(payload: dict = Body(...)):
    """Asigna manualmente un concepto de gasto en texto libre (ej. 'SUPER')
    a una categoria fija, o IGNORAR si no debe contar (nunca se adivina)."""
    concepto = payload.get("concepto")
    categoria_o_ignorar = payload.get("categoria_o_ignorar")
    if not concepto or not categoria_o_ignorar:
        raise HTTPException(status_code=400, detail="Falta 'concepto' o 'categoria_o_ignorar'.")
    dinero.guardar_mapeo_gasto(concepto, categoria_o_ignorar)
    return {"ok": True}


@app.get("/api/dinero/colchon")
async def ver_colchon(sucursal: str, mes: str):
    """mes en formato 'YYYY-MM'."""
    return colchon.calcular_debo_tener(sucursal, mes)


@app.post("/api/dinero/colchon/inicial")
async def guardar_colchon_inicial(payload: dict = Body(...)):
    """Colchon inicial del mes (el sobrante real del mes anterior) -- se
    captura a mano, nunca se asume $0."""
    sucursal = payload.get("sucursal")
    mes = payload.get("mes")
    monto = payload.get("monto")
    if not sucursal or not mes or monto is None:
        raise HTTPException(status_code=400, detail="Falta 'sucursal', 'mes' o 'monto'.")
    colchon.guardar_colchon_inicial(sucursal, mes, float(monto))
    return {"ok": True}


@app.post("/api/dinero/colchon/salida")
async def agregar_salida_efectivo(payload: dict = Body(...)):
    """Registra una salida grande de efectivo (nomina, compra al mayoreo) --
    usa las mismas categorias que los gastos chicos del corte diario."""
    fecha = payload.get("fecha")
    sucursal = payload.get("sucursal")
    concepto = payload.get("concepto")
    categoria = payload.get("categoria")
    monto = payload.get("monto")
    if not fecha or not sucursal or not concepto or not categoria or monto is None:
        raise HTTPException(status_code=400, detail="Falta algun campo obligatorio.")
    nuevo_id = colchon.agregar_salida_efectivo(fecha, sucursal, concepto, categoria, float(monto))
    return {"ok": True, "id": nuevo_id}


@app.delete("/api/dinero/colchon/salida/{id_salida}")
async def eliminar_salida_efectivo(id_salida: int):
    colchon.eliminar_salida_efectivo(id_salida)
    return {"ok": True}


@app.get("/api/compras")
async def listar_compras(
    ingrediente: str | None = None, proveedor: str | None = None,
    desde: str | None = None, hasta: str | None = None, sucursal: str | None = None,
):
    return compras.listar_compras(ingrediente=ingrediente, proveedor=proveedor, desde=desde, hasta=hasta, sucursal=sucursal)


@app.post("/api/compras")
async def agregar_compra(payload: dict = Body(...)):
    fecha = payload.get("fecha")
    ingrediente = payload.get("ingrediente")
    proveedor = payload.get("proveedor")
    cantidad = payload.get("cantidad")
    unidad = payload.get("unidad")
    precio_total = payload.get("precio_total")
    notas = payload.get("notas", "")
    sucursal = payload.get("sucursal")
    if not fecha or not ingrediente or not proveedor or precio_total is None:
        raise HTTPException(status_code=400, detail="Falta fecha, ingrediente, proveedor o precio.")
    try:
        nuevo_id = compras.agregar_compra(
            fecha, ingrediente, proveedor,
            float(cantidad) if cantidad not in (None, "") else None,
            unidad or None, float(precio_total), notas, sucursal or None,
        )
    except ValueError as e:
        raise HTTPException(status_code=400, detail=str(e))
    return {"ok": True, "id": nuevo_id}


@app.delete("/api/compras/{id_compra}")
async def eliminar_compra(id_compra: int):
    compras.eliminar_compra(id_compra)
    return {"ok": True}


@app.get("/api/compras/comparar")
async def comparar_precio_compra(ingrediente: str, precio_total: float, cantidad: float | None = None, unidad: str | None = None):
    return compras.comparar_precio(ingrediente, cantidad, unidad, precio_total)


@app.get("/api/compras/asesor/disponible")
async def asesor_disponible():
    return {"disponible": asesor_compras.disponible()}


@app.post("/api/compras/asesor")
async def preguntar_asesor(payload: dict = Body(...)):
    pregunta = payload.get("pregunta")
    ingrediente = payload.get("ingrediente")
    precio_total = payload.get("precio_total")
    cantidad = payload.get("cantidad")
    unidad = payload.get("unidad")
    if not pregunta or not ingrediente or precio_total is None:
        raise HTTPException(status_code=400, detail="Falta la pregunta, el ingrediente o el precio.")
    respuesta = asesor_compras.preguntar(pregunta, ingrediente, cantidad, unidad, float(precio_total))
    if respuesta is None:
        raise HTTPException(status_code=503, detail="El asesor con IA no está configurado (falta ANTHROPIC_API_KEY).")
    return {"respuesta": respuesta}


@app.get("/")
async def index():
    return FileResponse(STATIC_DIR / "index.html")


app.mount("/static", StaticFiles(directory=STATIC_DIR), name="static")
