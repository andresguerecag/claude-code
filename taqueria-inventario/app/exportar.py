"""Genera un Excel descargable a partir del reporte de conciliacion, para
que se pueda guardar/imprimir/mandar igual que hacian con sus reportes de
siempre."""
from __future__ import annotations

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

# Colores de marca T-Grill (tomados del logo)
NARANJA_OSCURO = "E8511C"
NARANJA = "F5821F"
AMARILLO = "FFCC00"
NEGRO = "1A1A1A"

FUENTE = "Arial"


def _encabezado(ws, fila, columnas, ancho_col=None):
    for i, texto in enumerate(columnas, start=1):
        c = ws.cell(row=fila, column=i, value=texto)
        c.font = Font(name=FUENTE, bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=NARANJA_OSCURO)
        c.alignment = Alignment(horizontal="center", vertical="center")
    if ancho_col:
        for i, ancho in enumerate(ancho_col, start=1):
            ws.column_dimensions[get_column_letter(i)].width = ancho


def generar_excel_reporte(reporte: dict) -> BytesIO:
    wb = Workbook()

    # --- Hoja 1: Comparativo ---
    ws = wb.active
    ws.title = "Comparativo"
    ws["A1"] = f"T-Grill — Conciliación de inventario — día {reporte.get('dia_detectado', '')}"
    ws["A1"].font = Font(name=FUENTE, bold=True, size=14, color=NARANJA_OSCURO)
    ws.merge_cells("A1:E1")

    if reporte.get("advertencia_confiabilidad"):
        ws["A2"] = "⚠ " + reporte["advertencia_confiabilidad"]
        ws["A2"].font = Font(name=FUENTE, italic=True, color="946200")
        ws["A2"].fill = PatternFill("solid", fgColor="FFF3CD")
        ws.merge_cells("A2:E2")
        ws.row_dimensions[2].height = 45
        ws["A2"].alignment = Alignment(wrap_text=True, vertical="top")

    fila = 4
    _encabezado(ws, fila, ["Insumo", "Consumo real", "Consumo teórico", "Diferencia", "Alerta"], ancho_col=[16, 15, 16, 14, 10])
    for f in reporte["comparativo"]:
        fila += 1
        ws.cell(row=fila, column=1, value=f["insumo"].capitalize())
        ws.cell(row=fila, column=2, value=f["consumo_real"])
        ws.cell(row=fila, column=3, value=f["consumo_teorico"])
        ws.cell(row=fila, column=4, value=f["diferencia"])
        ws.cell(row=fila, column=5, value="⚠ SI" if f["alerta"] else ("-" if f["alerta"] is False else "n/d"))
        if f["alerta"]:
            for col in range(1, 6):
                ws.cell(row=fila, column=col).fill = PatternFill("solid", fgColor="FFE5E0")
        for col in range(1, 6):
            ws.cell(row=fila, column=col).font = Font(name=FUENTE)

    # --- Hoja 2: Pendientes ---
    ws2 = wb.create_sheet("Pendientes de mapear")
    _encabezado(ws2, 1, ["Clave Wansoft", "Nombre", "Cantidad vendida"], ancho_col=[22, 40, 16])
    fila = 1
    for p in reporte["platillos_no_identificados"]:
        fila += 1
        ws2.cell(row=fila, column=1, value=p["clave"]).font = Font(name=FUENTE)
        ws2.cell(row=fila, column=2, value=p["nombre"]).font = Font(name=FUENTE)
        ws2.cell(row=fila, column=3, value=p["cantidad"]).font = Font(name=FUENTE)

    # --- Hoja 3: Inventario del día ---
    ws3 = wb.create_sheet("Inventario del día")
    _encabezado(ws3, 1, ["Producto", "Unidad", "Inicial", "Entrada", "Final", "Consumo real"], ancho_col=[22, 10, 12, 12, 12, 14])
    fila = 1
    for item in reporte["inventario_crudo"]:
        fila += 1
        ws3.cell(row=fila, column=1, value=item["producto"]).font = Font(name=FUENTE)
        ws3.cell(row=fila, column=2, value=item["unidad"]).font = Font(name=FUENTE)
        ws3.cell(row=fila, column=3, value=item["inicial"]).font = Font(name=FUENTE)
        ws3.cell(row=fila, column=4, value=item["entrada"]).font = Font(name=FUENTE)
        ws3.cell(row=fila, column=5, value=item["final"]).font = Font(name=FUENTE)
        ws3.cell(row=fila, column=6, value=item["consumo_real"]).font = Font(name=FUENTE)

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer
